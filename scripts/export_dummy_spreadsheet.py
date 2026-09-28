#!/usr/bin/env python3
"""
Skrip Ekspor Database Dummy Mock Segara ke Format Spreadsheet (.xlsx) dan CSV.

Output:
  1. Data/Database_Stok_Segara_Dummy.xlsx (Workbook Google Sheets lengkap)
  2. Data/csv_dummy/*.csv (File CSV per sheet)
"""

import os
import csv
import json
import subprocess
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.fonts import Font as StylesFont

StylesFont.family.max = 100

def main():
    print("Mengekstrak data dari mock backend Segara...")
    node_script = """
import { createMock } from './src/lib/mock.ts';
const mock = createMock();
const admin = mock.adminData('12345');
const history = mock.getLoginHistory('12345', 100);

const rekapRows = [];
const rekapBarisRows = [];
for (const r of admin.rekap) {
  rekapRows.push({
    id: r.id,
    ts: r.ts,
    waktu: r.waktu,
    karyawan_id: r.karyawan_id,
    karyawan: r.karyawan,
    diedit_admin: r.diedit_admin ?? false
  });
  if (r.baris) {
    for (const b of r.baris) {
      rekapBarisRows.push({
        rekap_id: b.rekap_id || r.id,
        barang_id: b.barang_id,
        barang: b.barang,
        saldo_awal: b.saldo_awal,
        diambil: b.diambil,
        sisa: b.sisa,
        terpakai: b.terpakai,
        catatan: b.catatan || ''
      });
    }
  }
}

const pengaturan = [
  { kunci: 'jam_tutup', nilai: admin.jamTutup || '21:00' },
  { kunci: 'urutan_kategori', nilai: JSON.stringify(admin.urutan || []) }
];

const karyawan = admin.karyawan.map(k => ({
  id: k.id,
  nama: k.nama,
  aktif: k.aktif,
  pin: k.nama === 'Budi Santoso' ? '1234' : ''
}));

const result = {
  Barang: admin.barang,
  Karyawan: karyawan,
  Transaksi: admin.transaksi,
  Rekap: rekapRows,
  RekapBaris: rekapBarisRows,
  Opname: admin.opname || [],
  Pengaturan: pengaturan,
  Log_Login: history || []
};

console.log(JSON.stringify(result));
"""
    res = subprocess.run(["node", "--experimental-strip-types", "-e", node_script], capture_output=True, text=True)
    if res.returncode != 0:
        raise RuntimeError(f"Gagal menjalankan node script: {res.stderr}")

    data = json.loads(res.stdout)

    SHEETS_DEF = {
        'Barang': ['id', 'nama', 'satuan', 'kategori', 'stok_dalam', 'stok_luar', 'ambang_min', 'alur', 'aktif', 'kode', 'catatan'],
        'Karyawan': ['id', 'nama', 'aktif', 'pin'],
        'Transaksi': ['id', 'ts', 'waktu', 'jenis', 'barang_id', 'barang', 'jumlah', 'karyawan_id', 'karyawan', 'alur', 'supplier', 'status', 'dicatat_oleh', 'catatan', 'kategori', 'satuan'],
        'Rekap': ['id', 'ts', 'waktu', 'karyawan_id', 'karyawan', 'diedit_admin'],
        'RekapBaris': ['rekap_id', 'barang_id', 'barang', 'saldo_awal', 'diambil', 'sisa', 'terpakai', 'catatan'],
        'Opname': ['id', 'ts', 'waktu', 'barang_id', 'barang', 'sistem', 'fisik', 'selisih'],
        'Pengaturan': ['kunci', 'nilai'],
        'Log_Login': ['id', 'ts', 'waktu', 'email', 'metode', 'role', 'status', 'user_agent'],
    }

    OUT_EXCEL = "Data/Database_Stok_Segara_Dummy.xlsx"
    CSV_DIR = "Data/csv_dummy"
    os.makedirs(CSV_DIR, exist_ok=True)

    print(f"Menyusun spreadsheet Excel {OUT_EXCEL}...")
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=10)
    border_thin = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    for sheet_name, headers in SHEETS_DEF.items():
        ws = wb_out.create_sheet(title=sheet_name)
        ws.append(headers)
        ws.freeze_panes = "A2"

        rows = data.get(sheet_name, [])

        # Header styling
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows
        for r_idx, row in enumerate(rows, start=2):
            vals = [row.get(h, '') for h in headers]
            ws.append(vals)
            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(r_idx, c_idx)
                c.font = cell_font
                c.border = border_thin

        # Column width
        for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for c in col:
                val_str = str(c.value or '')
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

        # CSV export
        csv_file = os.path.join(CSV_DIR, f"{sheet_name}.csv")
        with open(csv_file, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=headers, extrasaction='ignore')
            writer.writeheader()
            for row in rows:
                clean_row = {h: row.get(h, '') for h in headers}
                writer.writerow(clean_row)

    wb_out.save(OUT_EXCEL)
    print("Selesai! File berhasil dibuat:")
    print(f" - Excel: {OUT_EXCEL}")
    print(f" - CSV  : {CSV_DIR}/*.csv")

if __name__ == '__main__':
    main()
