#!/usr/bin/env python3
"""
Skrip Konversi Data Laporan Stock Barang Kedai Segara ke Format Sistem Stok Segara.

Input : Data/Laporan Stock Barang Kedai Segara.xlsx
Output: 
  1. Data/Database_Stok_Segara.xlsx (Workbook Google Sheets Segara lengkap)
  2. Data/csv/*.csv (CSV per sheet untuk kemudahan impor manual)
  3. apps-script/data_segara_september.js (Definisi DATA_SEGARA terbaru untuk Kode.gs)
"""

import os
import re
import csv
import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.fonts import Font as StylesFont

# Patch openpyxl untuk mengatasi file Excel dengan font family > 14
StylesFont.family.max = 100

SRC_FILE = "Data/Laporan Stock Barang Kedai Segara.xlsx"
OUT_EXCEL = "Data/Database_Stok_Segara.xlsx"
CSV_DIR = "Data/csv"
OUT_JS = "apps-script/data_segara_september.js"

CATEGORY_MAP = {
    'Freezer Protein': 'Freezer Protein',
    'Frezeer Bumbu': 'Freezer Bumbu',
    'Frezeer Roti & Juice': 'Freezer Roti & Juice',
    'Flavourful Drink': 'Flavourful Drink',
    'Stock Barang Kering (Dairy + Plant Base Milk)': 'Barang Kering (Dairy + Plant Base Milk)',
    'Kering (Dairy + Plant Base Milk)': 'Barang Kering (Dairy + Plant Base Milk)',
    'Bahan-Bahan Dasar + Kecap': 'Bahan Dasar + Kecap',
    'Bahan-Bahan Snack/ Dessert + Tepung': 'Bahan Snack / Dessert + Tepung',
    'Stock Cleaning Supplies + Utensils': 'Cleaning Supplies + Utensils',
    'Cleaning Supplies + Utensils': 'Cleaning Supplies + Utensils'
}

CATEGORY_ORDER = [
    'Freezer Protein',
    'Freezer Bumbu',
    'Freezer Roti & Juice',
    'Flavourful Drink',
    'Barang Kering (Dairy + Plant Base Milk)',
    'Bahan Dasar + Kecap',
    'Bahan Snack / Dessert + Tepung',
    'Cleaning Supplies + Utensils'
]

NOTES_PRESET = {
    'Ps.ukS': '1 Pack Isi 10 pcs',
    'Ps.ukM': '1 Pack Isi 10 pcs',
    'Ps.ukL': '1 Pack Isi 10 pcs',
    'Ps.ukXL': '1 Pack Isi 10 pcs',
}

def to_num(val):
    if val is None:
        return 0
    try:
        n = float(str(val).strip().replace(',', '.'))
        return int(n) if n.is_integer() else n
    except (ValueError, TypeError):
        return 0

def parse_date_val(val):
    if isinstance(val, (datetime.datetime, datetime.date)):
        return datetime.date(val.year, val.month, val.day)
    if isinstance(val, str):
        s = val.strip()
        m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
        if m:
            return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        m2 = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', s)
        if m2:
            return datetime.date(int(m2.group(3)), int(m2.group(2)), int(m2.group(1)))
    return None

def main():
    print(f"Membuka file sumber: {SRC_FILE}...")
    wb_src = openpyxl.load_workbook(SRC_FILE, data_only=True)
    ws_rekap = wb_src['Rekap Stock Barang']
    ws_masuk = wb_src['Barang Masuk ']
    ws_keluar = wb_src['Barang keluar']

    # 1. Parse Rekap Items
    items = []
    current_cat = None

    for r in range(1, ws_rekap.max_row + 1):
        c1 = ws_rekap.cell(r, 1).value
        c2 = ws_rekap.cell(r, 2).value
        if c1 is None:
            continue
        c1_str = str(c1).strip()
        if not c1_str:
            continue
        if c1_str.upper().startswith("LAPORAN") or c1_str.upper().startswith("CV.") or c1_str.upper().startswith("PERIODE:"):
            continue
        if c1_str == "Kode Barang":
            continue

        if c2 is None or str(c2).strip() == "":
            cat_raw = c1_str.replace("Stock Barang ", "").replace("Stock ", "").strip()
            current_cat = CATEGORY_MAP.get(cat_raw, cat_raw)
            continue

        kode = c1_str
        nama = str(c2).strip()
        satuan = str(ws_rekap.cell(r, 3).value or '').strip()

        stok_awal = to_num(ws_rekap.cell(r, 4).value)
        stok_masuk = to_num(ws_rekap.cell(r, 5).value)
        stok_keluar = to_num(ws_rekap.cell(r, 6).value)
        stok_akhir = to_num(ws_rekap.cell(r, 7).value)

        items.append({
            'kode': kode,
            'nama': nama,
            'satuan': satuan,
            'kategori': current_cat,
            'stok_awal': stok_awal,
            'stok_masuk': stok_masuk,
            'stok_keluar': stok_keluar,
            'stok_akhir': stok_akhir,
            'catatan': NOTES_PRESET.get(kode, '')
        })

    print(f"Berhasil membaca {len(items)} barang dari Rekap Stock Barang.")

    # Tambahkan BU jika belum ada (agar riwayat keluar tanggal 5 Sept terhubung)
    if not any(it['kode'].upper() == 'BU' for it in items):
        items.append({
            'kode': 'BU',
            'nama': 'Bumbu Ungkep',
            'satuan': 'Pack',
            'kategori': 'Freezer Bumbu',
            'stok_awal': 2,
            'stok_masuk': 0,
            'stok_keluar': 2,
            'stok_akhir': 0,
            'catatan': 'Menu lama / diganti Bumbu Ungkep Lidah (BUL)',
            'legacy': True
        })
        print("Menambahkan referensi legacy 'BU' (Bumbu Ungkep) untuk integritas relasi transaksi.")

    # Susun Barang
    barang_rows = []
    barang_by_code = {}
    for idx, it in enumerate(items, start=1):
        b_id = f"b{idx:03d}"
        it['id'] = b_id
        barang_by_code[it['kode'].upper()] = it

        stok_dalam = it['stok_akhir']
        stok_luar = 0
        ambang_min = 0
        alur = 'LUAR'
        aktif = False if it.get('legacy') else True
        opname_rekap = False if it.get('legacy') else True

        barang_rows.append({
            'id': b_id,
            'nama': it['nama'],
            'satuan': it['satuan'],
            'kategori': it['kategori'],
            'stok_dalam': stok_dalam,
            'stok_luar': stok_luar,
            'ambang_min': ambang_min,
            'alur': alur,
            'aktif': aktif,
            'kode': it['kode'],
            'catatan': it.get('catatan', ''),
            'opname_rekap': opname_rekap,
            '_awal': it['stok_awal'],
            '_masuk': it['stok_masuk'],
            '_keluar': it['stok_keluar'],
            '_akhir': it['stok_akhir']
        })

    # 2. Parse Transaksi Masuk & Keluar
    transaksi_rows = []

    # Masuk
    current_date = datetime.date(2026, 9, 1)
    daily_idx = 0
    for r in range(6, ws_masuk.max_row + 1):
        c1 = ws_masuk.cell(r, 1).value
        c2 = ws_masuk.cell(r, 2).value
        c3 = ws_masuk.cell(r, 3).value
        c4 = ws_masuk.cell(r, 4).value

        if c1 is None and c2 is None and c3 is None and c4 is None:
            continue

        d = parse_date_val(c1)
        note = ''
        if d is not None:
            if d != current_date:
                current_date = d
                daily_idx = 0
        elif c1 is not None and str(c1).strip():
            note = str(c1).strip()

        if c2 is not None and str(c2).strip():
            daily_idx += 1
            kode = str(c2).strip()
            nama = str(c3).strip() if c3 else ''
            qty = to_num(c4)

            b_info = barang_by_code.get(kode.upper())
            if b_info:
                b_id = b_info['id']
                b_name = b_info['nama']
                b_cat = b_info['kategori']
                b_unit = b_info['satuan']
            else:
                b_id = ''
                b_name = nama
                b_cat = ''
                b_unit = ''

            minute = (daily_idx - 1) % 60
            hour = 8 + (daily_idx - 1) // 60
            dt_wita = datetime.datetime(current_date.year, current_date.month, current_date.day, hour, minute)
            dt_utc = dt_wita - datetime.timedelta(hours=8)
            ts = int(dt_utc.replace(tzinfo=datetime.timezone.utc).timestamp() * 1000)
            waktu_str = dt_wita.strftime("%d/%m/%Y %H:%M")

            transaksi_rows.append({
                'ts': ts,
                'waktu': waktu_str,
                'jenis': 'MASUK',
                'barang_id': b_id,
                'barang': b_name,
                'jumlah': qty,
                'karyawan_id': 'k1',
                'karyawan': 'Admin Segara',
                'alur': 'DALAM',
                'supplier': 'CV. Dapur Rumah Rasa',
                'status': 'AKTIF',
                'dicatat_oleh': 'admin',
                'catatan': note,
                'kategori': b_cat,
                'satuan': b_unit,
            })

    # Keluar
    current_date = datetime.date(2026, 9, 1)
    daily_idx = 0
    for r in range(6, ws_keluar.max_row + 1):
        c1 = ws_keluar.cell(r, 1).value
        c2 = ws_keluar.cell(r, 2).value
        c3 = ws_keluar.cell(r, 3).value
        c4 = ws_keluar.cell(r, 4).value

        if c1 is None and c2 is None and c3 is None and c4 is None:
            continue

        d = parse_date_val(c1)
        note = ''
        if d is not None:
            if d != current_date:
                current_date = d
                daily_idx = 0
        elif c1 is not None and str(c1).strip():
            note = str(c1).strip()

        if c2 is not None and str(c2).strip():
            daily_idx += 1
            kode = str(c2).strip()
            nama = str(c3).strip() if c3 else ''
            qty = to_num(c4)

            b_info = barang_by_code.get(kode.upper())
            if b_info:
                b_id = b_info['id']
                b_name = b_info['nama']
                b_cat = b_info['kategori']
                b_unit = b_info['satuan']
            else:
                b_id = ''
                b_name = nama
                b_cat = ''
                b_unit = ''

            minute = (daily_idx - 1) % 60
            hour = 20 + (daily_idx - 1) // 60
            dt_wita = datetime.datetime(current_date.year, current_date.month, current_date.day, hour, minute)
            dt_utc = dt_wita - datetime.timedelta(hours=8)
            ts = int(dt_utc.replace(tzinfo=datetime.timezone.utc).timestamp() * 1000)
            waktu_str = dt_wita.strftime("%d/%m/%Y %H:%M")

            transaksi_rows.append({
                'ts': ts,
                'waktu': waktu_str,
                'jenis': 'AMBIL',
                'barang_id': b_id,
                'barang': b_name,
                'jumlah': qty,
                'karyawan_id': 'k2',
                'karyawan': 'Dapur Segara',
                'alur': 'LUAR',
                'supplier': '',
                'status': 'AKTIF',
                'dicatat_oleh': 'karyawan',
                'catatan': note,
                'kategori': b_cat,
                'satuan': b_unit,
            })

    # Urutkan transaksi secara kronologis
    transaksi_rows.sort(key=lambda x: x['ts'])
    for idx, tx in enumerate(transaksi_rows, start=1):
        tx['id'] = f"t{idx:04d}"

    print(f"Total transaksi berhasil diproses: {len(transaksi_rows)}")

    # 3. Karyawan
    karyawan_rows = [
        {'id': 'k1', 'nama': 'Admin Segara', 'aktif': True, 'pin': '12345'},
        {'id': 'k2', 'nama': 'Dapur Segara', 'aktif': True, 'pin': '1111'},
        {'id': 'k3', 'nama': 'Bar Segara', 'aktif': True, 'pin': '2222'},
    ]

    # 4. Pengaturan
    pengaturan_rows = [
        {'kunci': 'jam_tutup', 'nilai': '21:00'},
        {'kunci': 'urutan_kategori', 'nilai': json.dumps(CATEGORY_ORDER)},
    ]

    # 5. Buat Workbook Excel Hasil
    print(f"Menyusun file Excel baru: {OUT_EXCEL}...")
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)  # Hapus sheet default

    # Styling helper
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=10)
    border_thin = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    def write_sheet(name, headers, rows_data):
        ws = wb_out.create_sheet(title=name)
        ws.append(headers)
        ws.freeze_panes = "A2"

        # Header style
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Rows
        for r_idx, row in enumerate(rows_data, start=2):
            vals = [row.get(h, '') for h in headers]
            ws.append(vals)
            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(r_idx, c_idx)
                c.font = cell_font
                c.border = border_thin

        # Auto-adjust column width
        for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for c in col:
                val_str = str(c.value or '')
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    # 8 Sheet Segara Standar
    SHEETS_DEF = {
        'Barang': ['id', 'nama', 'satuan', 'kategori', 'stok_dalam', 'stok_luar', 'ambang_min', 'alur', 'aktif', 'kode', 'catatan', 'opname_rekap'],
        'Karyawan': ['id', 'nama', 'aktif', 'pin'],
        'Transaksi': ['id', 'ts', 'waktu', 'jenis', 'barang_id', 'barang', 'jumlah', 'karyawan_id', 'karyawan', 'alur', 'supplier', 'status', 'dicatat_oleh', 'catatan', 'kategori', 'satuan'],
        'Rekap': ['id', 'ts', 'waktu', 'karyawan_id', 'karyawan', 'diedit_admin'],
        'RekapBaris': ['rekap_id', 'barang_id', 'barang', 'saldo_awal', 'diambil', 'sisa', 'terpakai', 'catatan'],
        'Opname': ['id', 'ts', 'waktu', 'barang_id', 'barang', 'sistem', 'fisik', 'selisih'],
        'Pengaturan': ['kunci', 'nilai'],
        'Log_Login': ['id', 'ts', 'waktu', 'email', 'metode', 'role', 'status', 'user_agent'],
    }

    write_sheet('Barang', SHEETS_DEF['Barang'], barang_rows)
    write_sheet('Karyawan', SHEETS_DEF['Karyawan'], karyawan_rows)
    write_sheet('Transaksi', SHEETS_DEF['Transaksi'], transaksi_rows)
    write_sheet('Rekap', SHEETS_DEF['Rekap'], [])
    write_sheet('RekapBaris', SHEETS_DEF['RekapBaris'], [])
    write_sheet('Opname', SHEETS_DEF['Opname'], [])
    write_sheet('Pengaturan', SHEETS_DEF['Pengaturan'], pengaturan_rows)
    write_sheet('Log_Login', SHEETS_DEF['Log_Login'], [])

    # Sheet Tambahan: Ringkasan Rekap September untuk Referensi Cepat
    rekap_september_rows = []
    for b in barang_rows:
        rekap_september_rows.append({
            'kode': b['kode'],
            'nama': b['nama'],
            'satuan': b['satuan'],
            'kategori': b['kategori'],
            'stok_awal': b['_awal'],
            'stok_masuk': b['_masuk'],
            'stok_keluar': b['_keluar'],
            'stok_akhir': b['_akhir'],
            'status': 'Aktif' if b['aktif'] else 'Nonaktif (Legacy)'
        })
    write_sheet('Rekap_September', ['kode', 'nama', 'satuan', 'kategori', 'stok_awal', 'stok_masuk', 'stok_keluar', 'stok_akhir', 'status'], rekap_september_rows)

    wb_out.save(OUT_EXCEL)
    print(f"File Excel berhasil disimpan ke: {OUT_EXCEL}")

    # 6. Ekspor CSV untuk kemudahan impor manual ke Google Sheets
    print(f"Mengekspor file CSV ke folder: {CSV_DIR}...")
    os.makedirs(CSV_DIR, exist_ok=True)

    csv_targets = [
        ('Barang.csv', SHEETS_DEF['Barang'], barang_rows),
        ('Transaksi.csv', SHEETS_DEF['Transaksi'], transaksi_rows),
        ('Karyawan.csv', SHEETS_DEF['Karyawan'], karyawan_rows),
        ('Pengaturan.csv', SHEETS_DEF['Pengaturan'], pengaturan_rows),
        ('Rekap.csv', SHEETS_DEF['Rekap'], []),
        ('RekapBaris.csv', SHEETS_DEF['RekapBaris'], []),
        ('Opname.csv', SHEETS_DEF['Opname'], []),
        ('Log_Login.csv', SHEETS_DEF['Log_Login'], []),
        ('Rekap_September.csv', ['kode', 'nama', 'satuan', 'kategori', 'stok_awal', 'stok_masuk', 'stok_keluar', 'stok_akhir', 'status'], rekap_september_rows),
    ]

    for fname, cols, data in csv_targets:
        fpath = os.path.join(CSV_DIR, fname)
        with open(fpath, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(cols)
            for row in data:
                writer.writerow([row.get(c, '') for c in cols])

    print("Semua file CSV berhasil diekspor.")

    # 7. Generate array DATA_SEGARA terbaru untuk Kode.gs
    grouped = {}
    for cat in CATEGORY_ORDER:
        grouped[cat] = []

    for b in barang_rows:
        if b.get('aktif'):
            row_tuple = [b['kode'], b['nama'], b['satuan']]
            if b['catatan']:
                row_tuple.append(b['catatan'])
            cat = b['kategori']
            if cat not in grouped:
                grouped[cat] = []
            grouped[cat].append(row_tuple)

    js_lines = ["/* ---------- Data Kedai Segara (Laporan Stock September · CV. Dapur Rumah Rasa) ---------- */", "var DATA_SEGARA = ["]
    for cat in CATEGORY_ORDER:
        itms = grouped.get(cat, [])
        js_lines.append(f"  ['{cat}', [")
        for it in itms:
            js_lines.append(f"    {json.dumps(it)},")
        js_lines.append("  ]],")
    js_lines.append("];")

    with open(OUT_JS, 'w', encoding='utf-8') as f:
        f.write("\n".join(js_lines) + "\n")

    print(f"File JavaScript DATA_SEGARA tersimpan di: {OUT_JS}")
    print("\nPROSES KONVERSI SELESAI DENGAN SUKSES!")

if __name__ == '__main__':
    main()
